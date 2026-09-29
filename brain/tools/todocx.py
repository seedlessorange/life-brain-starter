#!/usr/bin/env python3
"""Markdown into a Word file that reads like a document.

    python3 brain/tools/todocx.py notes.md            # notes.docx beside it
    python3 brain/tools/todocx.py notes.md out.docx

Word does not understand markdown, so a .md pasted into a .docx arrives as
one long paragraph: the bullets lose their line breaks and the headings turn
into ordinary sentences. That is what happened to the Session 1 minutes. This
walks the markdown and builds real Word structure instead — headings, bullet
and numbered lists (nested), tables, quotes, code, bold and italic.

python-docx usually lives in the conda env rather than on the server's PATH,
so this re-runs itself there when it has to.
"""
import os
import re
import shutil
import subprocess
import sys

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:                                      # pragma: no cover
    for cand in ("~/miniconda3/bin/python", "~/anaconda3/bin/python"):
        exe = os.path.expanduser(cand)
        if os.path.exists(exe) and exe != sys.executable:
            os.execv(exe, [exe, os.path.abspath(__file__)] + sys.argv[1:])
    raise SystemExit("python-docx isn't installed for this python")

FENCE = re.compile(r"^(```|~~~)")
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
RULE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*\s][^*]*?\*|_[^_\s][^_]*?_|`[^`]+`)")


def _cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _runs(par, text):
    """Bold, italic and code inside a line, as real Word runs."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith(("**", "__")) and len(piece) > 4:
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Menlo"
            r.font.size = Pt(9.5)
        elif (piece.startswith(("*", "_")) and piece.endswith(("*", "_"))
              and len(piece) > 2):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def _style_for(kind, level):
    if kind == "bullet":
        return "List Bullet" if level == 0 else "List Bullet %d" % min(level + 1, 3)
    return "List Number" if level == 0 else "List Number %d" % min(level + 1, 3)


def convert(md_path, out_path=""):
    md_path = os.path.abspath(os.path.expanduser(md_path))
    out_path = os.path.abspath(os.path.expanduser(out_path)) if out_path else \
        os.path.splitext(md_path)[0] + ".docx"
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().replace("\r", "").split("\n")

    d = docx.Document()
    for section in d.sections:                    # a document, not a letter
        section.left_margin = section.right_margin = Pt(60)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        fence = FENCE.match(line)
        if fence:
            i += 1
            buf = []
            while i < len(lines) and not FENCE.match(lines[i].strip()):
                buf.append(lines[i])
                i += 1
            i += 1
            p = d.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Menlo"
            r.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(18)
            continue

        h = HEADING.match(line)
        if h:
            d.add_heading(h.group(2).strip(), level=min(len(h.group(1)), 4))
            i += 1
            continue

        if RULE.match(line):
            p = d.add_paragraph()
            p.add_run("").add_break()
            i += 1
            continue

        if "|" in line and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            head = _cells(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(_cells(lines[i]))
                i += 1
            t = d.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            for c, text in zip(t.rows[0].cells, head):
                c.paragraphs[0].add_run(text).bold = True
            for row in rows:
                cells = t.add_row().cells
                for c, text in zip(cells, row):
                    _runs(c.paragraphs[0], text)
            d.add_paragraph()
            continue

        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _runs(p, " ".join(buf))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        m = BULLET.match(raw) or NUMBER.match(raw)
        if m:
            kind = "bullet" if BULLET.match(raw) else "number"
            level = len(m.group(1).replace("\t", "    ")) // 2
            try:
                p = d.add_paragraph(style=_style_for(kind, level))
            except KeyError:
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Pt(18 * (level + 1))
            _runs(p, m.group(2).strip())
            i += 1
            continue

        buf = [line]
        i += 1
        while (i < len(lines) and lines[i].strip()
               and not HEADING.match(lines[i].strip())
               and not BULLET.match(lines[i]) and not NUMBER.match(lines[i])
               and not FENCE.match(lines[i].strip())
               and not lines[i].strip().startswith(">")):
            buf.append(lines[i].strip())
            i += 1
        p = d.add_paragraph()
        _runs(p, " ".join(buf))

    d.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    out = convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "")
    print("wrote", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
